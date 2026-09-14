# -*- coding: utf-8 -*-
"""
產生圖表 (PNG) 與彙整報告 Excel (combined_raw.xlsx)

五個指標家族 x 三種角度（皆為「該回合」的原始值，不是累積）：
  1. 總分 (score)         -- 各組平均 / 各組每回合平均 / 組內每位參與者每回合
  2. 擊中敵機 (hits)       -- 各組人類vsAI平均 / 各組每回合人類vsAI平均 / 組內每位參與者人類vsAI平均
  3. 準確率 (accuracy)     -- 各組人類vsAI / 各組每回合人類vsAI / 組內每位參與者人類vsAI
  4. 衝突次數 (conflict)   -- 各組平均 / 各組每回合平均 / 組內每位參與者每回合
  5. 訊號發送 (signals)    -- 各組人類vsAI / 各組每回合人類vsAI平均 / 組內每位參與者每回合人類vsAI
外加：
  6. 訊號內容拆解 -- 各組人類/AI 發送的訊號到底是「我來（自己接手）」還是「換你（交給對方）」
  7. 訊號後移動反應 -- 發送訊號後，發送者自己／接收方是靠近敵機、停留原地、還是遠離敵機
  8. 最後一次訊號 vs 該輪敵機最終由誰處理 -- 該回合最後一次訊號的語意，是否對得上該架
     敵機最後實際的結果（由誰擊落 / 被放過）
"""
import os
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

from analyze import (
    load_all, group_summary, round_group_summary, add_round_num, compute_hits,
    compute_conflicts_dedup, compute_signal_types, compute_signal_movement_response,
    summarize_movement_response, compute_last_signal_outcome, summarize_last_signal_outcome,
    GROUP_ORDER, GROUP_COLOR, ROLE_LABEL, OUT_DIR, CHART_DIR, CONFLICT_MERGE_GAP_SEC,
    MOVEMENT_WINDOW_SEC, MOVEMENT_STAY_THRESHOLD_PX,
)

# ---------------------------------------------------------------------------
# 中文字型
# ---------------------------------------------------------------------------
for _p in [
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
    "/usr/share/fonts/opentype/noto/NotoSansCJK-Bold.ttc",
]:
    if os.path.exists(_p):
        fm.fontManager.addfont(_p)

FONT_CANDIDATES = [
    "Noto Sans CJK TC", "Noto Sans CJK JP", "Noto Sans TC", "PingFang TC",
    "Heiti TC", "Microsoft JhengHei", "Arial Unicode MS",
]
_available = {f.name for f in fm.fontManager.ttflist}
_chosen = next((f for f in FONT_CANDIDATES if f in _available), None)
if _chosen:
    plt.rcParams["font.family"] = [_chosen]
plt.rcParams["axes.unicode_minus"] = False

# ---------------------------------------------------------------------------
# 色彩 / 樣式（dataviz 色票：類別色固定順序；人類 vs AI 用固定雙色，不隨組別變動）
# ---------------------------------------------------------------------------
SURFACE = "#fcfcfb"
INK_PRIMARY = "#0b0b0b"
INK_SECONDARY = "#52514e"
INK_MUTED = "#898781"
GRID = "#e1e0d9"
BASELINE = "#c3c2b7"

HUMAN_COLOR = "#2a78d6"   # blue  -- 全報告中「人類」固定用此色
AGENT_COLOR = "#eda100"   # amber -- 全報告中「AI」固定用此色

# 訊號後移動反應三分類固定配色（不隨組別 / 人類-AI 配色重複使用，避免混淆）
MOVE_COLOR = {
    "靠近／追擊敵機": "#1baf7a",   # 綠 -- 積極接手
    "停留原地":       "#c3c2b7",   # 灰 -- 沒有明顯動作
    "遠離敵機":       "#c0392b",   # 紅 -- 明確撤退／讓開
}
MOVE_ORDER = ["靠近／追擊敵機", "停留原地", "遠離敵機"]

# 最後一次訊號 vs 實際處理結果，三分類固定配色
OUTCOME_COLOR = {
    "預測命中": "#1baf7a",   # 綠 -- 訊號說了算，跟實際處理者一致
    "預測不符": "#c0392b",   # 紅 -- 實際處理者跟訊號說的不一樣
    "敵機被放過": "#898781",  # 灰 -- 沒有人擊落（missed）
}
OUTCOME_ORDER = ["預測命中", "預測不符", "敵機被放過"]

plt.rcParams.update({
    "figure.facecolor": SURFACE,
    "axes.facecolor": SURFACE,
    "axes.edgecolor": BASELINE,
    "axes.labelcolor": INK_SECONDARY,
    "text.color": INK_PRIMARY,
    "xtick.color": INK_MUTED,
    "ytick.color": INK_MUTED,
    "grid.color": GRID,
    "font.size": 11,
})


def _style_ax(ax, title=None, small=False):
    if title:
        ax.set_title(title, color=INK_PRIMARY, fontsize=11 if small else 13,
                      fontweight="bold", loc="left", pad=10)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_visible(False)
    ax.yaxis.grid(True, linewidth=0.8, color=GRID)
    ax.set_axisbelow(True)
    ax.tick_params(axis="both", length=0, labelsize=9 if small else 10)


def _savefig(fig, fname):
    path = os.path.join(CHART_DIR, fname)
    fig.savefig(path, facecolor=SURFACE, dpi=160)
    plt.close(fig)
    return path


def _suptitle(fig, text):
    fig.suptitle(text, x=0.02, y=0.99, ha="left", fontsize=14, fontweight="bold", color=INK_PRIMARY)


# ---------------------------------------------------------------------------
# 圖表 A：各組平均（單一數值，如總分 / 衝突次數）－ 長條圖 + SD 誤差線
# ---------------------------------------------------------------------------
def chart_group_bar_single(df, metric, title, ylabel, fname, id_col="experiment_id"):
    means, sds, colors, labels = [], [], [], []
    for g in GROUP_ORDER:
        vals = df.loc[df["group"] == g, metric].dropna().astype(float)
        means.append(vals.mean() if len(vals) else 0)
        sds.append(vals.std(ddof=1) if len(vals) > 1 else 0)
        colors.append(GROUP_COLOR[g])
        labels.append(g)

    fig, ax = plt.subplots(figsize=(7, 4.2))
    x = np.arange(len(labels))
    ax.bar(x, means, yerr=sds, capsize=4, color=colors, width=0.56,
           error_kw={"ecolor": INK_MUTED, "elinewidth": 1.2})
    for i, m in enumerate(means):
        ax.text(i, m + (max(means) * 0.03 if max(means) else 0.02), f"{m:.2f}",
                 ha="center", va="bottom", color=INK_SECONDARY, fontsize=10)
    ax.set_xticks(x)
    ax.set_xticklabels(labels, fontsize=10)
    ax.set_ylabel(ylabel)
    _style_ax(ax, title)
    fig.tight_layout()
    return _savefig(fig, fname)


# ---------------------------------------------------------------------------
# 圖表 B：各組人類 vs AI（雙數值）－ 分組長條圖
# ---------------------------------------------------------------------------
def chart_group_bar_dual(df, metric_a, metric_b, label_a, label_b, title, ylabel, fname):
    means_a, means_b = [], []
    for g in GROUP_ORDER:
        va = df.loc[df["group"] == g, metric_a].dropna().astype(float)
        vb = df.loc[df["group"] == g, metric_b].dropna().astype(float)
        means_a.append(va.mean() if len(va) else 0)
        means_b.append(vb.mean() if len(vb) else 0)

    fig, ax = plt.subplots(figsize=(7.5, 4.2))
    x = np.arange(len(GROUP_ORDER))
    w = 0.32
    ax.bar(x - w / 2, means_a, width=w, color=HUMAN_COLOR, label=label_a)
    ax.bar(x + w / 2, means_b, width=w, color=AGENT_COLOR, label=label_b)
    ax.set_xticks(x)
    ax.set_xticklabels(GROUP_ORDER, fontsize=10)
    ax.set_ylabel(ylabel)
    ax.legend(frameon=False, loc="upper right", fontsize=10)
    _style_ax(ax, title)
    fig.tight_layout()
    return _savefig(fig, fname)


# ---------------------------------------------------------------------------
# 圖表 C：各組每回合平均值（單一數值，非累積，該回合的原始平均）－ 一組一條線，同一張圖
# ---------------------------------------------------------------------------
def chart_group_round_line(round_all, metric, title, ylabel, fname):
    fig, ax = plt.subplots(figsize=(7.5, 4.5))
    for g in GROUP_ORDER:
        sub = round_all[round_all["group"] == g]
        agg = sub.groupby("round_num")[metric].mean().sort_index()
        ax.plot(agg.index, agg.values, marker="o", markersize=5, linewidth=2,
                color=GROUP_COLOR[g], label=g)
    ax.set_xlabel("回合 (Round)")
    ax.set_ylabel(ylabel)
    ax.legend(frameon=False, loc="best", fontsize=9)
    _style_ax(ax, title)
    fig.tight_layout()
    return _savefig(fig, fname)


# ---------------------------------------------------------------------------
# 圖表 D：組內每位參與者每回合數值（單一數值，非累積）－ 2x2 小圖，每小圖多條細線＝多位參與者
# ---------------------------------------------------------------------------
def chart_participant_round_grid(round_all, metric, suptitle, ylabel, fname):
    fig, axes = plt.subplots(2, 2, figsize=(10, 7.5), sharex=True, sharey=True)
    for ax, g in zip(axes.flat, GROUP_ORDER):
        sub = round_all[round_all["group"] == g].sort_values(["experiment_id", "round_num"])
        n_participants = sub["experiment_id"].nunique()
        for exp_id, part in sub.groupby("experiment_id"):
            ax.plot(part["round_num"], part[metric], linewidth=1.4, alpha=0.75, color=GROUP_COLOR[g])
        _style_ax(ax, f"{g}（n={n_participants}）", small=True)
        ax.set_xticks(sorted(round_all["round_num"].unique()))
    for ax in axes[-1, :]:
        ax.set_xlabel("回合 (Round)", fontsize=9)
    for ax in axes[:, 0]:
        ax.set_ylabel(ylabel, fontsize=9)
    _suptitle(fig, suptitle)
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    return _savefig(fig, fname)


# ---------------------------------------------------------------------------
# 圖表 E：各組每回合人類 vs AI 平均（非累積）－ 2x2 小圖，每小圖 2 條線
# ---------------------------------------------------------------------------
def chart_round_avg_dual_grid(round_all, metric_a, metric_b, label_a, label_b, suptitle, ylabel, fname):
    fig, axes = plt.subplots(2, 2, figsize=(10, 7.5), sharex=True, sharey=True)
    for ax, g in zip(axes.flat, GROUP_ORDER):
        sub = round_all[round_all["group"] == g]
        agg_a = sub.groupby("round_num")[metric_a].mean().sort_index()
        agg_b = sub.groupby("round_num")[metric_b].mean().sort_index()
        ax.plot(agg_a.index, agg_a.values, marker="o", markersize=4, linewidth=2, color=HUMAN_COLOR, label=label_a)
        ax.plot(agg_b.index, agg_b.values, marker="o", markersize=4, linewidth=2, color=AGENT_COLOR, label=label_b)
        _style_ax(ax, g, small=True)
        ax.set_xticks(sorted(round_all["round_num"].unique()))
    axes.flat[0].legend(frameon=False, loc="upper left", fontsize=8)
    for ax in axes[-1, :]:
        ax.set_xlabel("回合 (Round)", fontsize=9)
    for ax in axes[:, 0]:
        ax.set_ylabel(ylabel, fontsize=9)
    _suptitle(fig, suptitle)
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    return _savefig(fig, fname)


# ---------------------------------------------------------------------------
# 圖表 F：組內每位參與者人類 vs AI（session 總量）－ 2x2 小圖，每小圖分組長條
# ---------------------------------------------------------------------------
def chart_participant_bar_dual_grid(df, metric_a, metric_b, label_a, label_b, suptitle, ylabel, fname):
    fig, axes = plt.subplots(2, 2, figsize=(10, 7.5), sharey=True)
    for ax, g in zip(axes.flat, GROUP_ORDER):
        sub = df[df["group"] == g].sort_values("experiment_id")
        ids = sub["experiment_id"].astype(str).tolist()
        x = np.arange(len(ids))
        w = 0.35
        ax.bar(x - w / 2, sub[metric_a].values, width=w, color=HUMAN_COLOR, label=label_a)
        ax.bar(x + w / 2, sub[metric_b].values, width=w, color=AGENT_COLOR, label=label_b)
        ax.set_xticks(x)
        ax.set_xticklabels(ids, fontsize=8)
        ax.set_xlabel("參與者編號", fontsize=9)
        _style_ax(ax, g, small=True)
    axes.flat[0].legend(frameon=False, loc="upper right", fontsize=8)
    for ax in axes[:, 0]:
        ax.set_ylabel(ylabel, fontsize=9)
    _suptitle(fig, suptitle)
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    return _savefig(fig, fname)


# ---------------------------------------------------------------------------
# 圖表 G：組內每位參與者每回合人類 vs AI（非累積）－ 2x2 小圖，每人兩條線（實線=人類/虛線=AI）
# ---------------------------------------------------------------------------
def chart_participant_round_dual_grid(round_all, metric_a, metric_b, label_a, label_b, suptitle, ylabel, fname):
    fig, axes = plt.subplots(2, 2, figsize=(10, 7.5), sharex=True, sharey=True)
    for ax, g in zip(axes.flat, GROUP_ORDER):
        sub = round_all[round_all["group"] == g].sort_values(["experiment_id", "round_num"])
        n_participants = sub["experiment_id"].nunique()
        for exp_id, part in sub.groupby("experiment_id"):
            ax.plot(part["round_num"], part[metric_a], linewidth=1.2, alpha=0.7,
                     color=HUMAN_COLOR, linestyle="-")
            ax.plot(part["round_num"], part[metric_b], linewidth=1.2, alpha=0.7,
                     color=AGENT_COLOR, linestyle="--")
        _style_ax(ax, f"{g}（n={n_participants}）", small=True)
        ax.set_xticks(sorted(round_all["round_num"].unique()))
    from matplotlib.lines import Line2D
    handles = [
        Line2D([0], [0], color=HUMAN_COLOR, lw=2, label=label_a),
        Line2D([0], [0], color=AGENT_COLOR, lw=2, linestyle="--", label=label_b),
    ]
    axes.flat[0].legend(handles=handles, frameon=False, loc="upper left", fontsize=8)
    for ax in axes[-1, :]:
        ax.set_xlabel("回合 (Round)", fontsize=9)
    for ax in axes[:, 0]:
        ax.set_ylabel(ylabel, fontsize=9)
    _suptitle(fig, suptitle)
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    return _savefig(fig, fname)


# ---------------------------------------------------------------------------
# 圖表 H：各組人類 / AI 發送的訊號內容拆解 － 2x2 小圖，堆疊長條（自己接手 vs 交給對方）
# ---------------------------------------------------------------------------
def chart_signal_type_breakdown(signal_types_df, suptitle, ylabel, fname):
    fig, axes = plt.subplots(2, 2, figsize=(10, 7.5), sharey=True)
    for ax, g in zip(axes.flat, GROUP_ORDER):
        sub = signal_types_df[signal_types_df["group"] == g]
        senders = ["人類", "AI"]
        colors = [HUMAN_COLOR, AGENT_COLOR]
        i_can_vals, your_turn_vals = [], []
        for s in senders:
            row = sub[sub["sender"] == s]
            i_can_vals.append(row.loc[row["signal_type"] == "i_can", "count"].sum())
            your_turn_vals.append(row.loc[row["signal_type"] == "your_turn", "count"].sum())

        x = np.arange(len(senders))
        w = 0.5
        bars1 = ax.bar(x, i_can_vals, width=w, color=colors, alpha=0.55, label="我來（i_can，自己接手）")
        bars2 = ax.bar(x, your_turn_vals, width=w, bottom=i_can_vals, color=colors, alpha=1.0,
                        label="換你（your_turn，交給對方）")
        for xi, (a, b) in enumerate(zip(i_can_vals, your_turn_vals)):
            if a:
                ax.text(xi, a / 2, str(int(a)), ha="center", va="center", fontsize=8, color="white")
            if b:
                ax.text(xi, a + b / 2, str(int(b)), ha="center", va="center", fontsize=8, color="white")
        ax.set_xticks(x)
        ax.set_xticklabels(senders, fontsize=10)
        _style_ax(ax, g, small=True)

    from matplotlib.patches import Patch
    handles = [
        Patch(facecolor=INK_MUTED, alpha=0.55, label="我來（自己接手，不交接）"),
        Patch(facecolor=INK_MUTED, alpha=1.0, label="換你（交給對方）"),
    ]
    axes.flat[0].legend(handles=handles, frameon=False, loc="upper right", fontsize=7.5)
    for ax in axes[:, 0]:
        ax.set_ylabel(ylabel, fontsize=9)
    _suptitle(fig, suptitle)
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    return _savefig(fig, fname)


# ---------------------------------------------------------------------------
# 圖表 I：訊號後移動反應 -- 2x2 小圖（依組別），x 軸為「角色 x 訊號類型」
# （人類-我來／人類-換你／AI-我來／AI-換你），100% 堆疊長條
# （靠近／追擊敵機 vs 停留原地 vs 遠離敵機），可分別畫「發送者自己」或「接收方」
# 角色（人類 / AI）明確拆開，不會把兩者的反應混在一起平均。
# ---------------------------------------------------------------------------
def chart_movement_stack_grid(movement_summary, perspective, suptitle, fname):
    fig, axes = plt.subplots(2, 2, figsize=(11, 8), sharey=True)
    signal_labels = {"i_can": "我來", "your_turn": "換你"}
    role_order = ["human", "agent"]
    combos = [(r, st) for r in role_order for st in ["i_can", "your_turn"]]
    xlabels = [f"{ROLE_LABEL[r]}\n{signal_labels[st]}" for r, st in combos]
    for ax, g in zip(axes.flat, GROUP_ORDER):
        sub = movement_summary[(movement_summary["group"] == g) & (movement_summary["perspective"] == perspective)]
        x = np.arange(len(combos))
        bottoms = np.zeros(len(combos))
        totals = np.array(
            [sub[(sub["role"] == r) & (sub["signal_type"] == st)]["count"].sum() for r, st in combos],
            dtype=float,
        )
        for mv in MOVE_ORDER:
            vals = np.array([
                sub[(sub["role"] == r) & (sub["signal_type"] == st) & (sub["movement"] == mv)]["count"].sum()
                for r, st in combos
            ], dtype=float)
            pct = np.divide(vals * 100, totals, out=np.zeros_like(vals), where=totals > 0)
            ax.bar(x, pct, bottom=bottoms, width=0.6, color=MOVE_COLOR[mv], label=mv)
            for xi, (p, b, t) in enumerate(zip(pct, bottoms, totals)):
                if p >= 8 and t > 0:
                    ax.text(xi, b + p / 2, f"{p:.0f}%", ha="center", va="center", fontsize=7.5, color="white")
            bottoms += pct
        for xi, t in enumerate(totals):
            if t == 0:
                ax.text(xi, 3, "—", ha="center", va="bottom", fontsize=11, color=INK_MUTED)
        ax.set_xticks(x)
        ax.set_xticklabels(xlabels, fontsize=8)
        ax.set_ylim(0, 100)
        _style_ax(ax, g, small=True)
    from matplotlib.patches import Patch
    handles = [Patch(facecolor=MOVE_COLOR[m], label=m) for m in MOVE_ORDER]
    axes.flat[0].legend(handles=handles, frameon=True, framealpha=0.9, edgecolor="none",
                          facecolor=SURFACE, loc="upper right", fontsize=7.5)
    for ax in axes[:, 0]:
        ax.set_ylabel("佔比 (%)", fontsize=9)
    _suptitle(fig, suptitle)
    fig.tight_layout(rect=[0, 0, 1, 0.94])
    return _savefig(fig, fname)


# ---------------------------------------------------------------------------
# 圖表 J：各組「最後一次訊號」是否對得上該輪敵機最終處理結果 -- 100% 堆疊長條
# ---------------------------------------------------------------------------
def chart_last_signal_outcome(outcome_summary, title, ylabel, fname):
    fig, ax = plt.subplots(figsize=(7.5, 4.5))
    x = np.arange(len(GROUP_ORDER))
    bottoms = np.zeros(len(GROUP_ORDER))
    cat_vals = {}
    for g in GROUP_ORDER:
        row = outcome_summary[outcome_summary["group"] == g]
        if row.empty or pd.isna(row["match_rate_among_resolved"].values[0]):
            cat_vals[g] = {"預測命中": 0, "預測不符": 0, "敵機被放過": 0}
            continue
        r = row.iloc[0]
        n_total = r["n_rounds_with_signal"]
        n_match = r["n_predicted_match"]
        n_mismatch = r["n_resolved"] - r["n_predicted_match"]
        n_missed = r["n_missed"]
        cat_vals[g] = {
            "預測命中": n_match / n_total * 100 if n_total else 0,
            "預測不符": n_mismatch / n_total * 100 if n_total else 0,
            "敵機被放過": n_missed / n_total * 100 if n_total else 0,
        }
    for cat in OUTCOME_ORDER:
        vals = np.array([cat_vals[g][cat] for g in GROUP_ORDER])
        ax.bar(x, vals, bottom=bottoms, width=0.56, color=OUTCOME_COLOR[cat], label=cat)
        for xi, (v, b) in enumerate(zip(vals, bottoms)):
            if v >= 5:
                ax.text(xi, b + v / 2, f"{v:.0f}%", ha="center", va="center", fontsize=9, color="white")
        bottoms += vals
    ax.set_xticks(x)
    ax.set_xticklabels(GROUP_ORDER, fontsize=10)
    ax.set_ylabel(ylabel)
    ax.set_ylim(0, 100)
    ax.legend(frameon=False, loc="upper right", fontsize=9)
    _style_ax(ax, title)
    fig.tight_layout()
    return _savefig(fig, fname)


# ---------------------------------------------------------------------------
# Excel 輸出
# ---------------------------------------------------------------------------
HEADER_FILL = PatternFill("solid", fgColor="2a78d6")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def write_df_sheet(wb, name, df, freeze="A2"):
    ws = wb.create_sheet(name)
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for i, col in enumerate(df.columns, start=1):
        try:
            maxlen = max([len(str(col))] + [len(str(v)) for v in df[col].astype(str).values[:200]])
        except Exception:
            maxlen = 12
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, maxlen + 2), 40)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    return ws


def main():
    experiment_all, round_all, event_all = load_all()
    round_all = add_round_num(round_all)

    hits_by_exp, hits_by_round = compute_hits(event_all)
    # 併入 experiment 層級：total 擊中次數 + 準確率交叉核對
    experiment_wide = experiment_all.merge(hits_by_exp, on=["group", "experiment_id"], how="left")
    experiment_wide[["human_hits", "agent_hits"]] = experiment_wide[["human_hits", "agent_hits"]].fillna(0)
    # round 層級也併入 hits，供每回合平均圖使用
    round_wide = round_all.merge(hits_by_round, on=["group", "experiment_id", "round_num"], how="left")
    round_wide[["human_hits", "agent_hits"]] = round_wide[["human_hits", "agent_hits"]].fillna(0)

    # 衝突次數：新增「連續衝突只算一次」欄位（total_conflict_dedup / conflict_count_dedup），
    # 併入 wide 表，與原始 total_conflict / conflict_count 並列，方便比較。
    conflict_by_exp, conflict_by_round = compute_conflicts_dedup(event_all)
    experiment_wide = experiment_wide.merge(conflict_by_exp, on=["group", "experiment_id"], how="left")
    experiment_wide["total_conflict_dedup"] = experiment_wide["total_conflict_dedup"].fillna(0)
    round_wide = round_wide.merge(conflict_by_round, on=["group", "experiment_id", "round_num"], how="left")
    round_wide["conflict_count_dedup"] = round_wide["conflict_count_dedup"].fillna(0)

    # 統計摘要改用併入新欄位後的 wide 表，這樣 total_conflict_dedup / conflict_count_dedup
    # 也會一併算出各組平均、SD、Kruskal-Wallis 檢定，方便跟原始欄位並排比較。
    exp_summary = group_summary(experiment_wide)
    round_summary = round_group_summary(round_wide)

    signal_types = compute_signal_types(event_all)

    # 訊號後移動反應 + 最後一次訊號 vs 敵機最終處理結果
    movement_detail = compute_signal_movement_response(event_all)
    movement_summary = summarize_movement_response(movement_detail)
    last_signal_outcome = compute_last_signal_outcome(event_all)
    last_signal_summary = summarize_last_signal_outcome(last_signal_outcome)
    # Excel 不支援帶時區的 datetime，寫入前先去掉 tzinfo（僅顯示用，不影響計算）
    for _df in (movement_detail, last_signal_outcome):
        if "timestamp" in _df.columns and len(_df):
            _df["timestamp"] = pd.to_datetime(_df["timestamp"]).dt.tz_localize(None)

    charts = []

    # ---- 1. 總分 ----
    charts.append(chart_group_bar_single(
        experiment_all, "total_score", "各組平均總分 (± SD)", "總分", "01_score_group_avg.png"))
    charts.append(chart_group_round_line(
        round_wide, "round_score", "各組每回合平均總分", "平均總分／回合", "02_score_round_growth.png"))
    charts.append(chart_participant_round_grid(
        round_wide, "round_score", "組內每位參與者每回合總分", "總分", "03_score_participant_growth.png"))

    # ---- 2. 擊中敵機 ----
    charts.append(chart_group_bar_dual(
        experiment_wide, "human_hits", "agent_hits", "人類擊中數", "AI 擊中數",
        "各組人類 vs AI 擊中敵機平均數", "擊中次數（每人總計）", "04_hits_group_avg.png"))
    charts.append(chart_round_avg_dual_grid(
        round_wide, "human_hits", "agent_hits", "人類擊中數", "AI 擊中數",
        "各組每回合人類 vs AI 擊中敵機平均數", "平均擊中次數／回合", "05_hits_round_avg.png"))
    charts.append(chart_participant_bar_dual_grid(
        experiment_wide, "human_hits", "agent_hits", "人類擊中數", "AI 擊中數",
        "組內每位參與者人類 vs AI 擊中敵機數", "擊中次數（session 總計）", "06_hits_participant_avg.png"))

    # ---- 3. 準確率 ----
    charts.append(chart_group_bar_dual(
        experiment_all, "human_total_accuracy", "agent_total_accuracy", "人類準確率", "AI 準確率",
        "各組人類 vs AI 準確率", "準確率", "07_accuracy_group_avg.png"))
    charts.append(chart_round_avg_dual_grid(
        round_all, "human_accuracy", "agent_accuracy", "人類準確率", "AI 準確率",
        "各組每回合人類 vs AI 準確率", "平均準確率／回合", "08_accuracy_round_avg.png"))
    charts.append(chart_participant_bar_dual_grid(
        experiment_all, "human_total_accuracy", "agent_total_accuracy", "人類準確率", "AI 準確率",
        "組內每位參與者人類 vs AI 準確率", "準確率（session 總計）", "09_accuracy_participant_avg.png"))

    # ---- 4. 衝突次數 ----
    charts.append(chart_group_bar_single(
        experiment_all, "total_conflict", "各組平均衝突次數 (± SD)", "衝突次數", "10_conflict_group_avg.png"))
    charts.append(chart_group_round_line(
        round_all, "conflict_count", "各組每回合平均衝突次數", "平均衝突次數／回合", "11_conflict_round_growth.png"))
    charts.append(chart_participant_round_grid(
        round_all, "conflict_count", "組內每位參與者每回合衝突次數", "衝突次數", "12_conflict_participant_growth.png"))

    # ---- 4b. 衝突次數：原始 vs 連續衝突只算一次（新增比較，不刪除原有 10~12） ----
    charts.append(chart_group_bar_dual(
        experiment_wide, "total_conflict", "total_conflict_dedup", "原始（每次都算）", "去重複（連續衝突算一次）",
        f"各組平均衝突次數：原始 vs 去重複（間隔 <= {CONFLICT_MERGE_GAP_SEC:.0f}s 視為同一段）",
        "衝突次數", "17_conflict_compare_group_avg.png"))
    charts.append(chart_round_avg_dual_grid(
        round_wide, "conflict_count", "conflict_count_dedup", "原始（每回合）", "去重複（每回合）",
        "各組每回合衝突次數：原始 vs 去重複", "平均衝突次數／回合", "18_conflict_compare_round_avg.png"))
    charts.append(chart_participant_round_dual_grid(
        round_wide, "conflict_count", "conflict_count_dedup", "原始", "去重複",
        "組內每位參與者每回合衝突次數：原始 vs 去重複", "衝突次數", "19_conflict_compare_participant.png"))

    # ---- 5. 訊號發送次數 ----
    charts.append(chart_group_bar_dual(
        experiment_all, "human_total_signals", "agent_total_signals", "人類發出信號數", "AI 發出信號數",
        "各組人類 vs AI 訊號發送次數", "次數", "13_signals_group_avg.png"))
    charts.append(chart_round_avg_dual_grid(
        round_all, "human_signal_count", "agent_signal_count", "人類發出信號數", "AI 發出信號數",
        "各組每回合人類 vs AI 平均訊號發送次數", "平均次數／回合", "14_signals_round_growth.png"))
    charts.append(chart_participant_round_dual_grid(
        round_all, "human_signal_count", "agent_signal_count", "人類發出信號數", "AI 發出信號數",
        "組內每位參與者每回合人類 vs AI 訊號發送次數", "次數", "15_signals_participant_growth.png"))

    # ---- 6. 訊號內容拆解（交給人類 vs 交給AI vs 自己接手）----
    charts.append(chart_signal_type_breakdown(
        signal_types, "各組人類 / AI 發送訊號內容拆解（我來 vs 換你）", "訊號次數", "16_signal_type_breakdown.png"))

    # ---- 7. 訊號後移動反應（發送訊號後 3 秒內，發送者自己 vs 接收方是否真的照訊號行動）----
    charts.append(chart_movement_stack_grid(
        movement_summary, "發送者自己",
        f"發送訊號後 {MOVEMENT_WINDOW_SEC:.0f} 秒內：發送者自己的移動反應",
        "20_signal_movement_sender.png"))
    charts.append(chart_movement_stack_grid(
        movement_summary, "接收方（對方）",
        f"發送訊號後 {MOVEMENT_WINDOW_SEC:.0f} 秒內：接收方（對方）的移動反應",
        "21_signal_movement_receiver.png"))

    # ---- 8. 最後一次訊號 vs 該輪敵機最終由誰處理 ----
    charts.append(chart_last_signal_outcome(
        last_signal_summary, "各組：最後一次訊號的預測 vs 該輪敵機最終處理結果",
        "佔「有發送訊號」回合數的比例 (%)", "22_last_signal_vs_outcome.png"))

    # ---- Excel ----
    wb = Workbook()
    wb.remove(wb.active)

    write_df_sheet(wb, "Raw_Experiment", experiment_wide)
    write_df_sheet(wb, "Raw_Round", round_wide)
    write_df_sheet(wb, "Raw_Events", event_all)
    write_df_sheet(wb, "Group_Summary_Experiment", exp_summary)
    write_df_sheet(wb, "Group_Summary_Round", round_summary)
    write_df_sheet(wb, "Signal_Type_Breakdown", signal_types)
    write_df_sheet(wb, "Signal_Movement_Detail", movement_detail)
    write_df_sheet(wb, "Signal_Movement_Summary", movement_summary)
    write_df_sheet(wb, "Last_Signal_Outcome", last_signal_outcome)
    write_df_sheet(wb, "Last_Signal_Outcome_Summary", last_signal_summary)

    ws_charts = wb.create_sheet("Charts", 0)
    ws_charts.sheet_view.showGridLines = False
    ws_charts["A1"] = "Floating Rescue 實驗資料 — 四組比較圖表"
    ws_charts["A1"].font = Font(size=16, bold=True)
    ws_charts["A2"] = "資料來源：game/data/<組別>/<實驗編號>/{experiment,round,events}.csv　（重新執行 build_report.py 可更新）"
    ws_charts["A2"].font = Font(size=10, color="898781")

    row = 4
    for path in charts:
        img = XLImage(path)
        is_grid = any(k in path for k in ["participant", "round_avg", "round_growth", "breakdown", "movement"]) and "group_avg" not in path
        if is_grid:
            img.width, img.height = 640, 480
        else:
            img.width, img.height = 560, 320
        ws_charts.add_image(img, f"A{row}")
        row += 27 if is_grid else 18

    out_path = os.path.join(OUT_DIR, "combined_raw.xlsx")
    wb.save(out_path)
    print("Saved:", out_path)
    print("Charts dir:", CHART_DIR)
    return out_path


if __name__ == "__main__":
    main()
